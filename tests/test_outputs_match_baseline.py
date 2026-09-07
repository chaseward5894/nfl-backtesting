"""Single acceptance test: nflbt-run output must match the baseline reference files."""

import json
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
REFERENCE_ROOT = (
    REPO_ROOT / "NFL-Model-UPDATED" / "review2026" / "baseline"
)
PACKAGE_ROOT = REPO_ROOT / "src"

CASES = [
    {
        "name": "dataset_full",
        "dataset": "DATASET_FULL",
        "metrics_ref": REFERENCE_ROOT / "metrics.txt",
        "xlsx_ref": REFERENCE_ROOT / "baseline.xlsx",
    },
    {
        "name": "dataset_2025",
        "dataset": "DATASET_2025",
        "metrics_ref": (
            REFERENCE_ROOT / "backtesting_2025" / "metrics_2025.txt"
        ),
        "xlsx_ref": (
            REFERENCE_ROOT / "backtesting_2025" / "backtesting_2025.xlsx"
        ),
    },
]


def _run_cli(tmp_path: Path, dataset: str) -> Path:
    out = tmp_path / "out"
    out.mkdir()
    env = os.environ.copy()
    env["PYTHONPATH"] = str(PACKAGE_ROOT)
    result = subprocess.run(
        [sys.executable, "-m", "nflbetting_backtest.cli",
         dataset, "--out", str(out)],
        env=env,
        capture_output=True,
        text=True,
        cwd=str(REPO_ROOT),
    )
    assert result.returncode == 0, (
        f"CLI failed:\nstdout={result.stdout}\nstderr={result.stderr}"
    )
    # CLI creates a timestamped subdirectory under out. Find it.
    children = list(out.iterdir())
    assert len(children) == 1, (
        f"expected exactly one run subdirectory under {out}, got {children}"
    )
    return children[0]


@pytest.mark.parametrize("case", CASES, ids=[c["name"] for c in CASES])
def test_outputs_match_baseline(tmp_path, case):
    run_dir = _run_cli(tmp_path, case["dataset"])

    # All expected outputs are present
    expected_files = [
        "metrics.txt", "predictions.xlsx", "predictions.jsonl",
        "predictions.pkl", "config.yaml.copy", "run_card.json",
    ]
    for fname in expected_files:
        assert (run_dir / fname).exists(), f"missing output: {fname}"

    metrics_actual = (run_dir / "metrics.txt").read_text()
    metrics_expected = case["metrics_ref"].read_text()
    assert metrics_actual == metrics_expected, (
        f"metrics.txt for {case['name']} differs.\n"
        f"--- diff (head) ---\n"
        + "\n".join(
            difflib_unified_diff(
                metrics_expected.splitlines(),
                metrics_actual.splitlines(),
                lineterm="",
            )[:40]
        )
    )

    wb_actual = openpyxl.load_workbook(run_dir / "predictions.xlsx")
    wb_expected = openpyxl.load_workbook(case["xlsx_ref"])
    assert wb_actual.sheetnames == wb_expected.sheetnames, (
        f"sheet names differ: {wb_actual.sheetnames} vs {wb_expected.sheetnames}"
    )
    actual_rows = list(wb_actual.active.iter_rows(values_only=True))
    expected_rows = list(wb_expected.active.iter_rows(values_only=True))
    assert len(actual_rows) == len(expected_rows), (
        f"row count differs: {len(actual_rows)} vs {len(expected_rows)}"
    )
    for i, (a, e) in enumerate(zip(actual_rows, expected_rows)):
        for j, (av, ev) in enumerate(zip(a, e)):
            if isinstance(av, float) and isinstance(ev, float):
                assert abs(av - ev) <= 0.5, (
                    f"row {i} col {j}: {av} vs {ev}"
                )
            else:
                assert av == ev, (
                    f"row {i} col {j}: {av!r} vs {ev!r}"
                )

    # Run card sanity checks
    card = json.loads((run_dir / "run_card.json").read_text())
    assert card["dataset"] == case["dataset"]
    assert card["n_predictions"] == len(actual_rows) - 1  # minus header
    assert card["config"]["values"]["training_weeks_used"] > 0
    assert card["config"]["values"]["test_weeks_used"] > 0
    assert card["outputs"]["config_copy"] == "config.yaml.copy"
    assert Path(card["config"]["source_path"]).exists()


def difflib_unified_diff(a, b, lineterm=""):
    import difflib
    return list(difflib.unified_diff(a, b, lineterm=lineterm))